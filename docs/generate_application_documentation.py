"""Generate the source-derived ONERC Core application documentation.

Run from the app root with the bench Python environment::

    /home/nigel/frappe/main-bench/env/bin/python \
        apps/onerc_core/docs/generate_application_documentation.py

The generator deliberately reads DocType JSON at build time. Field and permission
appendices therefore describe the checked-out source rather than a hand-maintained
copy that can silently drift.
"""

from __future__ import annotations

import ast
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(__file__).resolve().parent / "ONERC_Core_Application_Documentation.docx"
AS_OF = date(2026, 8, 31)

RED = RGBColor(0xA6, 0x19, 0x2E)
DARK = RGBColor(0x20, 0x27, 0x2D)
GREY = RGBColor(0x5D, 0x65, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = "F3F5F7"
PALE_RED = "F8ECEF"


def shade(cell, fill: str):
	props = cell._tc.get_or_add_tcPr()
	element = OxmlElement("w:shd")
	element.set(qn("w:fill"), fill)
	props.append(element)


def set_cell_text(cell, text: object, bold=False, colour=None, size=8):
	cell.text = ""
	p = cell.paragraphs[0]
	r = p.add_run("" if text is None else str(text))
	r.bold = bold
	r.font.name = "Aptos"
	r.font.size = Pt(size)
	if colour:
		r.font.color.rgb = colour
	cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction: str):
	run = OxmlElement("w:r")
	begin = OxmlElement("w:fldChar")
	begin.set(qn("w:fldCharType"), "begin")
	text = OxmlElement("w:instrText")
	text.set(qn("xml:space"), "preserve")
	text.text = instruction
	separate = OxmlElement("w:fldChar")
	separate.set(qn("w:fldCharType"), "separate")
	end = OxmlElement("w:fldChar")
	end.set(qn("w:fldCharType"), "end")
	run.extend([begin, text, separate, end])
	paragraph._p.append(run)


class Report:
	def __init__(self):
		self.doc = Document()
		self._styles()
		self._page()

	def _styles(self):
		styles = self.doc.styles
		normal = styles["Normal"]
		normal.font.name = "Aptos"
		normal.font.size = Pt(9.5)
		normal.font.color.rgb = DARK
		normal.paragraph_format.space_after = Pt(5)
		normal.paragraph_format.line_spacing = 1.08
		for name, size, colour, before, after in (
			("Title", 34, RED, 0, 10),
			("Subtitle", 15, GREY, 0, 14),
			("Heading 1", 21, RED, 16, 7),
			("Heading 2", 14, DARK, 12, 5),
			("Heading 3", 11, RED, 9, 3),
		):
			style = styles[name]
			style.font.name = "Aptos Display"
			style.font.size = Pt(size)
			style.font.bold = True
			style.font.color.rgb = colour
			style.paragraph_format.space_before = Pt(before)
			style.paragraph_format.space_after = Pt(after)

	def _page(self):
		section = self.doc.sections[0]
		section.top_margin = Inches(0.65)
		section.bottom_margin = Inches(0.65)
		section.left_margin = Inches(0.7)
		section.right_margin = Inches(0.7)
		header = section.header.paragraphs[0]
		header.text = "ONERC CORE  /  APPLICATION REFERENCE"
		header.style = self.doc.styles["Caption"]
		header.runs[0].font.color.rgb = RED
		footer = section.footer.paragraphs[0]
		footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
		footer.add_run("Source-derived • 31 August 2026  |  ")
		add_field(footer, "PAGE")

	def title_page(self):
		p = self.doc.add_paragraph()
		p.paragraph_format.space_before = Inches(1.15)
		r = p.add_run("ONERC")
		r.bold = True
		r.font.name = "Aptos Display"
		r.font.size = Pt(22)
		r.font.color.rgb = RED
		p = self.doc.add_paragraph("Core Application Documentation", style="Title")
		p = self.doc.add_paragraph(
			"Functional, technical, data, security, integration, deployment and operations reference",
			style="Subtitle",
		)
		self.callout(
			"Document status",
			"Source-derived baseline for the checked-out ONERC Core repository as of 31 August 2026. "
			"Runtime configuration, installed satellite apps, custom fields, roles and production data are site-specific and are not inferred.",
		)
		self.table(
			["Attribute", "Value"],
			[
				["Application", "onerc_core / Onerc Core"],
				["Platform", "Frappe Framework 16.x"],
				["Python", ">= 3.14"],
				["License", "MIT"],
				["Publisher in hooks", "Kelvin Njenga"],
				["Document date", "31 August 2026"],
			],
		)
		self.doc.add_page_break()

	def h1(self, text):
		self.doc.add_heading(text, 1)

	def h2(self, text):
		self.doc.add_heading(text, 2)

	def h3(self, text):
		self.doc.add_heading(text, 3)

	def p(self, text, bold_lead=None):
		p = self.doc.add_paragraph()
		if bold_lead:
			p.add_run(bold_lead).bold = True
		p.add_run(text)
		return p

	def bullets(self, items):
		for item in items:
			self.doc.add_paragraph(item, style="List Bullet")

	def numbered(self, items):
		for item in items:
			self.doc.add_paragraph(item, style="List Number")

	def callout(self, title, text, colour=PALE_RED):
		t = self.doc.add_table(rows=1, cols=1)
		t.alignment = WD_TABLE_ALIGNMENT.CENTER
		cell = t.cell(0, 0)
		shade(cell, colour)
		p = cell.paragraphs[0]
		r = p.add_run(f"{title}\n")
		r.bold = True
		r.font.color.rgb = RED
		p.add_run(text)
		self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

	def table(self, headers, rows, widths=None):
		t = self.doc.add_table(rows=1, cols=len(headers))
		t.style = "Table Grid"
		t.alignment = WD_TABLE_ALIGNMENT.CENTER
		for i, header in enumerate(headers):
			shade(t.rows[0].cells[i], "A6192E")
			set_cell_text(t.rows[0].cells[i], header, True, WHITE, 8)
		for row_no, row in enumerate(rows):
			cells = t.add_row().cells
			for i, value in enumerate(row):
				if row_no % 2:
					shade(cells[i], LIGHT)
				set_cell_text(cells[i], value, size=7.6)
		if widths:
			for row in t.rows:
				for i, width in enumerate(widths):
					row.cells[i].width = Inches(width)
		self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
		return t

	def toc(self):
		self.h1("Contents")
		p = self.doc.add_paragraph()
		add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
		self.p("Open the document in Microsoft Word and choose Update Field to refresh page numbers.")
		self.doc.add_page_break()

	def save(self):
		props = self.doc.core_properties
		props.title = "ONERC Core Application Documentation"
		props.subject = "Functional and technical reference"
		props.author = "Generated from the ONERC Core source repository"
		props.keywords = "ONERC, Frappe, architecture, DocType, API, security"
		self.doc.save(OUTPUT)


def doctypes():
	items = []
	for path in sorted(ROOT.glob("onerc_core/**/doctype/*/*.json")):
		data = json.loads(path.read_text())
		if data.get("doctype") != "DocType":
			continue
		data["_path"] = path.relative_to(ROOT).as_posix()
		items.append(data)
	return items


def python_inventory():
	rows = []
	for path in sorted(ROOT.glob("onerc_core/**/*.py")):
		try:
			tree = ast.parse(path.read_text())
		except SyntaxError:
			continue
		public = []
		for node in tree.body:
			if isinstance(node, (ast.FunctionDef, ast.ClassDef)) and not node.name.startswith("_"):
				public.append(node.name)
		if public:
			rows.append([path.relative_to(ROOT).as_posix(), ", ".join(public)])
	return rows


def add_overview(r: Report, dts):
	r.h1("1. Executive overview")
	r.p(
		"ONERC Core is the shared foundation of a configurable Red Cross / Red Crescent management ecosystem. "
		"It supplies cross-product identity, geography, organisation settings, access scoping and reusable content services. "
		"Product-specific satellite applications can integrate through hooks without Core importing their domain models."
	)
	r.h2("Purpose and boundaries")
	r.bullets(
		[
			"Provide a thin, reusable person identity spine through Red Profile, while keeping volunteer, member, employee and beneficiary domain facts in satellite applications.",
			"Represent a society-defined geographic hierarchy using levels and nested-set nodes, with optional coordinates and compatibility adapters.",
			"Bind Frappe roles to geographic nodes using time-bound Geo Assignments and enforce the resulting scope consistently.",
			"Centralise society branding, locale, terminology, validation rules and feature switches.",
			"Provide reusable knowledge articles, FAQs, feedback intake, stakeholder records and related lookup vocabularies.",
			"Remain inert where an integration hook has no registrations; Core must not invent product-specific policy.",
		]
	)
	r.h2("Repository profile")
	tests = list(ROOT.glob("onerc_core/**/test_*.py")) + list(ROOT.glob("onerc_core/**/tests/test_*.py"))
	r.table(
		["Metric", "Observed in source"],
		[
			["DocTypes", len(dts)],
			["Child-table DocTypes", sum(bool(x.get("istable")) for x in dts)],
			["Single settings DocTypes", sum(bool(x.get("issingle")) for x in dts)],
			["Python files", len(list(ROOT.glob("onerc_core/**/*.py")))],
			["Test modules", len(set(tests))],
			["Seed fixture files", len(list(ROOT.glob("onerc_core/fixtures/*.json")))],
		],
	)
	r.h2("Primary audiences")
	r.table(
		["Audience", "How to use this document"],
		[
			[
				"Product owners and society administrators",
				"Understand capabilities, configuration choices and operational workflows.",
			],
			["Developers", "Use the architecture, service contracts, API catalogue and DocType dictionary."],
			[
				"Security and compliance reviewers",
				"Review trust boundaries, guest endpoints, personal data and geo-scope enforcement.",
			],
			[
				"Implementers and operators",
				"Install, configure, migrate, test, monitor and troubleshoot the app.",
			],
		],
	)


def add_architecture(r: Report):
	r.h1("2. Architecture and design")
	r.h2("Logical architecture")
	r.table(
		["Layer", "Responsibilities", "Representative components"],
		[
			[
				"Presentation / Desk",
				"Frappe forms, list/tree views and the public OneRC workspace.",
				"DocType JSON, client scripts, Geo Node tree script",
			],
			[
				"HTTP API",
				"Guest and authenticated endpoints for content, settings, FAQ and feedback use cases.",
				"onerc_core.api.* and whitelisted service methods",
			],
			[
				"Application services",
				"Policy-oriented interfaces shared by Core and satellite apps.",
				"access, geo, identity and society services",
			],
			[
				"Domain documents",
				"Validation, naming, reconciliation and lifecycle logic.",
				"Red Profile, Geo Node, Geo Assignment, Article, Settings",
			],
			[
				"Persistence",
				"Frappe DocTypes, child tables, Singles, nested-set indices and framework records.",
				"MariaDB/PostgreSQL through Frappe ORM",
			],
			[
				"Extension boundary",
				"Dependency-inverted hooks implemented by installed satellite apps.",
				"affiliation providers, capability resolver, scopeable doctypes",
			],
		],
	)
	r.h2("Dependency direction")
	r.callout(
		"Core rule",
		"Satellite applications may depend on Core. Core does not import product-specific satellite applications. Integrations are declared through Frappe hooks and resolved at runtime.",
	)
	r.table(
		["Hook", "Provided by", "Contract and default behaviour"],
		[
			[
				"onerc_affiliation_providers",
				"Satellite apps",
				"Provider owns declared reference DocTypes and returns live affiliation claims. No provider means rebuild is a no-op.",
			],
			[
				"onerc_capability_resolver",
				"Exactly one policy owner",
				"Answers whether a user holds a gated capability. No resolver means gated rows are hidden.",
			],
			[
				"onerc_scopeable_doctypes",
				"Owner of each product DocType",
				"Declares doctype, geo field, and exactly one role source. No registrations means scoping is inert.",
			],
		],
	)
	r.h2("Key design principles")
	r.bullets(
		[
			"Fail closed at access boundaries: missing assignment, unresolved role or missing capability does not widen access.",
			"Single source of truth: geographic traversal goes through the adapter; live authority comes from Geo Assignment; society UI configuration comes through the configuration service.",
			"Configurable society policy: names, colours, locales, phone validation, terminology, feature keys and product scope roles are not embedded in product code.",
			"Thin shared identity: personal facts and affiliation pointers are central; domain-specific operational records remain outside Core.",
			"Opaque identifiers: Geo Node names use a series so mutable labels and levels are not database keys.",
		]
	)


def add_modules(r: Report):
	r.h1("3. Functional modules")
	sections = [
		(
			"Identity and affiliation",
			[
				"Red Profile records a person's shared identity: name, contact details, optional demographics, residence, identifications, linked Frappe user and derived affiliations.",
				"The profile supports thin registration: personal completeness is intentionally not forced by the shared spine.",
				"Affiliation providers declare ownership of satellite reference DocTypes and return current claims; rebuild reconciles only provider-owned rows, preventing unrelated integrations from deleting one another's data.",
				"Affiliation Type may associate a read capability. The read gate removes sensitive affiliation rows unless the configured resolver grants that capability.",
			],
		),
		(
			"Geography",
			[
				"Geo Level defines ordered, configurable tiers and whether a parent is required. Geo Node implements a nested-set tree with parent-depth and sibling-name validation.",
				"Geo Node stores an optional latitude/longitude pair. A half-pair and coordinates outside earth bounds are rejected.",
				"The adapter exposes roots, children, ancestors, descendants, levels, paths, point lookup, upward resolution and scope matching without leaking tree-storage details to consumers.",
				"Region and selector DocTypes provide compatibility and UI selection structures around the canonical Geo Node model.",
			],
		),
		(
			"Access control and approver routing",
			[
				"Geo Assignment binds a user and a role to a Geo Node, with activation and optional validity dates. The assignment grants authority only while it is live and the user still holds the named Frappe role.",
				"Scope expansion includes the assigned node and all descendants. Multiple disjoint assignments are unioned; ancestors and siblings are not accidentally included.",
				"Three enforcement layers are supported: list-query conditions, record permission checks and an explicit API guard.",
				"System Manager and Administrator have an explicit unrestricted bypass. Society-specific authority remains configurable.",
				"Approver routing can select holders at a configured level or walk upward to the nearest ancestor with a live holder.",
			],
		),
		(
			"National society configuration",
			[
				"National Society Settings is a Single DocType for identity, branding, theme tokens, locale, social channels, terminology, feature toggles, validation policy and product-added scope-role settings.",
				"The UI configuration service applies field-level fallbacks: System Settings for locale, primary logo for dark logo, caller defaults for absent feature keys, and key text for absent terminology.",
				"Validation checks IANA time zones, compiles phone-number regular expressions, validates settings-backed Frappe roles, normalises keys to lower snake case and rejects duplicates.",
			],
		),
		(
			"Knowledge articles",
			[
				"Article supports draft/submitted lifecycle, article type, category, pillar, authorship, cover content, featured ordering, SEO metadata and engagement counters.",
				"Validation generates unique slugs, estimates reading time at 200 words per minute, validates future scheduled publication and supplies metadata defaults.",
				"A five-minute scheduler publishes due submitted articles. Public reads require Published status and normally docstatus 1.",
				"Authenticated users can toggle likes; views and engagement can be read publicly.",
			],
		),
		(
			"FAQ and feedback",
			[
				"Available FAQ categories and published FAQs are guest-readable. Search covers question and answer. Voting atomically increments helpful or not-helpful counters.",
				"Feedback types are publicly listed and feedback can be submitted by guests or authenticated users. Administrative workflows use Open, Reviewed, Resolved and Closed states.",
			],
		),
		(
			"Stakeholders and communication",
			[
				"Stakeholder captures organisation/consortium/individual records, classifications, interest areas, contacts and supporting details through child tables and controlled vocabularies.",
				"SMS Campaign provides campaign metadata and message/audience configuration. Its current Python controller contains no sending integration or orchestration logic; deployments should not assume delivery from this app alone.",
			],
		),
	]
	for title, paragraphs in sections:
		r.h2(title)
		for paragraph in paragraphs:
			r.p(paragraph)


def add_workflows(r: Report):
	r.h1("4. Core workflows")
	r.h2("Profile and affiliation lifecycle")
	r.numbered(
		[
			"Create a Red Profile with the minimum identity required by the invoking product; optional personal fields remain optional.",
			"Validate names, contact policy, residence shape, user uniqueness and identification consistency.",
			"A satellite application creates or updates its domain record and calls the affiliation service, or participates in a full rebuild through its provider hook.",
			"Core validates claims and reconciles provider-owned references into read-only Red Profile Affiliation rows.",
			"When the profile is read, the gate consults Affiliation Type capabilities and removes rows the current user may not see.",
		]
	)
	r.h2("Geo-scoped record access")
	r.numbered(
		[
			"The product app registers its DocType, Geo Node field and literal or settings-backed role in onerc_scopeable_doctypes.",
			"An administrator grants the Frappe role to a user and creates a live Geo Assignment at the intended node.",
			"The scope service verifies both the live assignment and current role membership, then expands each assigned subtree.",
			"List reads receive an SQL condition; direct record reads invoke has_permission; custom endpoints should invoke guard before returning or mutating a named record.",
			"Removing the role or expiring/deactivating the assignment removes authority on the next check; the service does not cache stale scopes.",
		]
	)
	r.h2("Scheduled article publication")
	r.numbered(
		[
			"An editor prepares an Article and selects Scheduled with a future Scheduled Publish At value.",
			"The article must be submitted (docstatus 1), forming the approval gate.",
			"Every five minutes the scheduler finds due submitted articles and changes status to Published while setting Published On.",
			"Failures are logged per article; a successful batch is committed. Draft and cancelled rows are ignored.",
		]
	)
	r.h2("Feedback handling")
	r.numbered(
		[
			"A client obtains configured Feedback Types.",
			"The visitor supplies a non-empty message, valid type and optional subject/email. Authenticated submissions record the session user.",
			"Authorised staff list records using pagination and optional status/type filters.",
			"Reviewers move a record to Reviewed, Resolved or Closed and capture reviewer identity, timestamp and notes.",
		]
	)


def add_api(r: Report):
	r.h1("5. API and service reference")
	r.p(
		"Frappe exposes whitelisted Python methods at /api/method/<dotted.path>. Authentication labels below reflect decorators in source; normal DocType permissions and explicit checks remain relevant."
	)
	rows = [
		[
			"GET/POST",
			"onerc_core.api.article.get_articles",
			"Guest",
			"Filters: article_type, category, pillar, is_featured, limit, offset, include_drafts; returns ordered article summaries.",
		],
		[
			"GET",
			"onerc_core.api.article.get_article",
			"Guest",
			"slug; returns submitted published article and increments view_count.",
		],
		[
			"POST",
			"onerc_core.api.article.create_article",
			"Authenticated",
			"Article fields; inserts through normal permissions.",
		],
		["GET", "onerc_core.api.article.get_categories", "Guest", "Returns active Localisation Categories."],
		[
			"POST",
			"onerc_core.api.article.toggle_like",
			"Authenticated",
			"article_slug; creates/deletes Article Like and updates counter.",
		],
		[
			"GET",
			"onerc_core.api.article.get_article_engagement",
			"Guest",
			"article_slug; returns like count, actual Comment count and current-user liked flag.",
		],
		[
			"GET",
			"onerc_core.api.faq.get_faq_categories",
			"Guest",
			"Returns available categories in display order.",
		],
		["GET", "onerc_core.api.faq.get_faq_list", "Guest", "category, search; returns published FAQs."],
		["GET", "onerc_core.api.faq.get_faq", "Guest", "name; rejects unpublished records."],
		["POST", "onerc_core.api.faq.vote_faq", "Guest", "name, vote; vote is helpful or not_helpful."],
		["GET", "onerc_core.api.feedback.get_feedback_types", "Guest", "Returns all feedback types."],
		[
			"POST",
			"onerc_core.api.feedback.submit_feedback",
			"Guest",
			"feedback_type, message, optional subject/email.",
		],
		[
			"GET",
			"onerc_core.api.feedback.get_feedback_list",
			"Authenticated",
			"status, feedback_type, page, page_size (1–100).",
		],
		[
			"POST",
			"onerc_core.api.feedback.update_feedback",
			"Authenticated + role",
			"name, Reviewed/Resolved/Closed, notes; only LH Admin or System Manager.",
		],
		[
			"POST",
			"onerc_core.api.organization.update_organization_settings",
			"Guest",
			"Updates name, short name and logo; see security review note.",
		],
		[
			"GET",
			"onerc_core.api.organization.get_organization_settings",
			"Guest",
			"Returns public branding/contact configuration.",
		],
		[
			"GET",
			"onerc_core.society.services.config.get_ui_config",
			"Authenticated",
			"Returns society, theme, locale, terminology, features, phone policy and social media.",
		],
		[
			"POST",
			"...doctype.article.article.increment_view_count",
			"Guest",
			"Article name; increments and returns view counter.",
		],
	]
	r.table(["Method", "Dotted endpoint", "Access", "Inputs / behaviour"], rows)
	r.h2("Internal service contracts")
	r.table(
		["Service", "Supported responsibilities"],
		[
			[
				"geo.services.adapter",
				"Tree navigation, configured level descriptions, point reads, upward resolution and matching.",
			],
			[
				"access.services.registry",
				"Read and validate scope registrations; resolve literal/settings-backed roles.",
			],
			[
				"access.services.scope",
				"Unrestricted check, live assignment nodes, expanded user scope and exact-node holders.",
			],
			[
				"access.services.enforcement",
				"Permission-query SQL, record permission, is_in_scope and explicit guard.",
			],
			["access.services.approvers", "Resolve exact-level or nearest-ancestor approvers."],
			["identity.services.affiliation", "Set and rebuild provider-owned Red Profile affiliations."],
			["identity.services.read_gate", "Capability lookup and redaction of gated affiliations."],
			["society.services.config", "Stable UI config payload, terms, features and phone validation."],
		],
	)


def add_security(r: Report):
	r.h1("6. Security, privacy and controls")
	r.h2("Implemented controls")
	r.bullets(
		[
			"Frappe role and DocType permissions provide the baseline authorization model.",
			"Wildcard permission hooks are inert for unregistered DocTypes and apply geo filtering only to declared product records.",
			"Geo scope fails closed when no assignment grants, a settings-backed role cannot resolve, or a registered geo field is invalid.",
			"Role removal immediately invalidates assignment authority because live checks read current Has Role rows.",
			"Capability-gated affiliations are hidden when no resolver is installed or access is denied.",
			"Article creation uses normal insert permissions; feedback update explicitly limits users to LH Admin or System Manager.",
			"Validation prevents malformed time zones, regular expressions, tree relationships, coordinates and duplicate configuration keys.",
		]
	)
	r.h2("Personal and potentially sensitive data")
	r.table(
		["Area", "Examples", "Recommended handling"],
		[
			[
				"Red Profile",
				"Names, email, phone, birth date, gender, marital status, citizenship, address, photo, IDs",
				"Apply least privilege, retention rules, audit review, encryption/backups policy and data-subject procedures.",
			],
			[
				"Affiliations",
				"Relationship type, status and satellite record pointer",
				"Configure capability gates for sensitive relationship types.",
			],
			[
				"Feedback",
				"Message, email, submitter and reviewer notes",
				"Avoid collecting unnecessary data; limit reviewer access and define retention.",
			],
			[
				"Geo Assignment",
				"User, role, organisational territory and dates",
				"Treat as authorization data; monitor changes and review periodically.",
			],
			[
				"Stakeholders",
				"Contacts, classifications and interests",
				"Confirm lawful basis, access, retention and export requirements.",
			],
		],
	)
	r.h2("Security review observations")
	r.callout(
		"High-priority review",
		"update_organization_settings is decorated allow_guest=True and can save organisation name, short name and logo. Unless a separate upstream control is guaranteed, change this endpoint to authenticated/authorised access before exposing the site publicly.",
	)
	r.bullets(
		[
			"get_feedback_list is authenticated but its source contains a commented-out role restriction. Confirm DocType read permissions are intentionally sufficient or add an explicit administrative role check.",
			"Guest feedback and FAQ voting endpoints should be protected by deployment-level rate limiting and abuse monitoring.",
			"Article and FAQ counters call explicit commits. Account for their write load, bot traffic and transaction semantics.",
			"The hooks.py user_data_fields section is not configured. Implement privacy export/deletion mappings if the deployment relies on Frappe personal-data tooling.",
			"SMS Campaign contains no provider integration in this repository. Any downstream sender must protect credentials, consent, suppression lists and delivery callbacks.",
		]
	)


def add_operations(r: Report):
	r.h1("7. Installation, configuration and operations")
	r.h2("Prerequisites")
	r.bullets(
		[
			"A functioning Frappe Bench with Frappe >=16.0.0 and <17.0.0.",
			"Python 3.14 or later as declared by the project.",
			"Database, Redis, workers, scheduler and web processes configured according to the Frappe deployment model.",
			"Access to the ONERC Core Git repository and any intended satellite applications.",
		]
	)
	r.h2("Installation")
	r.numbered(
		[
			"From the bench directory, obtain the application with bench get-app <repository-url> --branch <branch>.",
			"Install it on the target site with bench --site <site> install-app onerc_core.",
			"Run bench --site <site> migrate after upgrades so schema, fixtures and patches are applied.",
			"Build assets and restart processes using the deployment's normal Bench or process-manager workflow.",
		]
	)
	r.h2("Initial configuration checklist")
	r.bullets(
		[
			"Complete National Society Settings: four required identity fields are organisation name, short name, country and primary language; branding images are optional.",
			"Review theme, locale, currency, time zone, date/time formats and social links.",
			"Configure and test the phone regular expression with a representative example.",
			"Define terminology and feature keys expected by installed products; duplicate normalised keys are rejected.",
			"Create Geo Levels, then Geo Nodes from roots downward; verify hierarchy overview and optional map points.",
			"Create roles required by products, populate settings-backed role fields, grant roles to users and create matching Geo Assignments.",
			"Review Identification Types, Stakeholder Entities and Stakeholder Classifications seeded by fixtures; sites may deactivate or extend the vocabularies.",
			"Confirm scheduler operation so due articles publish on the five-minute cron interval.",
		]
	)
	r.h2("Seed data and migrations")
	r.p(
		"Fixtures seed four identification keys (national_id, passport, alien_id, driving_licence), three stakeholder entity names and six stakeholder classifications. Filters restrict export to the seeded vocabulary so society-added records are not swept into Core fixtures."
	)
	r.p(
		"patches.txt currently references the nationality-to-country-of-citizenship migration. Treat patches as forward-only deployment steps and back up the site before migration."
	)
	r.h2("Monitoring and troubleshooting")
	r.table(
		["Symptom", "Checks"],
		[
			[
				"User sees no scoped records",
				"Confirm registration, resolved role, current role membership, active/dates on Geo Assignment and record Geo Node.",
			],
			[
				"All users denied after configuration",
				"Inspect Error Log for unresolved scope-role title and validate National Society Settings role values.",
			],
			[
				"Scheduled article remains scheduled",
				"Confirm docstatus=1, due timestamp/time zone, scheduler enabled and worker/error logs.",
			],
			[
				"Affiliation missing",
				"Confirm provider hook, owned reference_doctypes, returned claim validity and rebuild result.",
			],
			[
				"Affiliation unexpectedly hidden",
				"Confirm Affiliation Type capability, exactly one resolver and resolver decision for the user.",
			],
			[
				"Phone rejected",
				"Inspect configured regex and example; remember full-match semantics are used.",
			],
			[
				"Geo tree save fails",
				"Check required parent, level order, duplicate sibling label and coordinate pair/ranges.",
			],
		],
	)


def add_testing(r: Report):
	r.h1("8. Testing and quality assurance")
	r.p(
		"The repository uses Frappe IntegrationTestCase tests. The strongest suites exercise access scope, registry validation, enforcement, approver routing, geography, affiliations, read gating, society configuration and Red Profile invariants. Several scaffolded DocType tests contain pass only, so file count must not be mistaken for comprehensive behavioural coverage."
	)
	r.h2("Recommended commands")
	r.table(
		["Purpose", "Command"],
		[
			["All app tests", "bench --site <site> run-tests --app onerc_core"],
			["One module", "bench --site <site> run-tests --module <python.test.module>"],
			["Lint/format checks", "pre-commit run --all-files"],
			["Migration rehearsal", "bench --site <staging-site> migrate"],
		],
	)
	r.h2("Release verification")
	r.bullets(
		[
			"Back up and restore a representative staging site before production migration.",
			"Run the full app suite and targeted tests for all modified domains.",
			"Verify guest endpoints, authenticated endpoints and explicit role checks with positive and negative users.",
			"Test geo access at assigned node, descendant, sibling, ancestor, expired assignment and revoked-role cases.",
			"Exercise provider absence, malformed provider response and capability-resolver absence.",
			"Verify content scheduling across the configured time zone and scheduler interval.",
			"Review generated schema diffs, fixture changes, patches and error logs before go-live.",
		]
	)
	r.h2("Coverage priorities")
	r.p(
		"Add meaningful tests for Article API permissions and counters, FAQ voting abuse controls, Feedback list authorization, guest organisation settings mutation, Stakeholder workflows and any eventual SMS provider. These areas are either public-facing, security-sensitive or currently represented by scaffold tests."
	)


def add_data_dictionary(r: Report, dts):
	r.h1("9. Complete DocType data dictionary")
	r.p(
		"This appendix is generated directly from every DocType JSON file in the repository. Layout-only fields are included so administrators and implementers can map the Desk form precisely. ‘Options’ contains Link targets, Select choices or table child DocTypes as applicable."
	)
	for dt in dts:
		r.h2(dt.get("name", "Unnamed DocType"))
		traits = []
		if dt.get("issingle"):
			traits.append("Single")
		if dt.get("istable"):
			traits.append("Child table")
		if dt.get("is_tree"):
			traits.append("Tree")
		if dt.get("is_submittable"):
			traits.append("Submittable")
		traits.append(f"Module: {dt.get('module', '—')}")
		traits.append(f"Naming: {dt.get('autoname') or 'controller/framework default'}")
		r.p(" • ".join(traits))
		rows = []
		for f in dt.get("fields", []):
			flags = []
			for key, label in (
				("reqd", "required"),
				("unique", "unique"),
				("read_only", "read-only"),
				("hidden", "hidden"),
				("in_list_view", "list"),
				("translatable", "translatable"),
			):
				if f.get(key):
					flags.append(label)
			depends = f.get("depends_on") or f.get("mandatory_depends_on") or f.get("read_only_depends_on")
			if depends:
				flags.append(f"conditional: {depends}")
			options = f.get("options", "")
			if isinstance(options, str) and "\n" in options:
				options = " | ".join(options.splitlines())
			rows.append(
				[
					f.get("label", ""),
					f.get("fieldname", ""),
					f.get("fieldtype", ""),
					options,
					", ".join(flags),
				]
			)
		if rows:
			r.table(["Label", "Field name", "Type", "Options / target", "Rules"], rows)
		else:
			r.p("No fields are declared in this DocType JSON.")
		permissions = dt.get("permissions", [])
		if permissions:
			perm_rows = []
			keys = (
				"read",
				"write",
				"create",
				"delete",
				"submit",
				"cancel",
				"amend",
				"report",
				"export",
				"share",
				"print",
				"email",
			)
			for p in permissions:
				grants = ", ".join(key for key in keys if p.get(key)) or "none"
				perm_rows.append([p.get("role", ""), p.get("permlevel", 0), grants])
			r.table(["Role", "Level", "Granted actions"], perm_rows)
		r.p(f"Source: {dt['_path']}")


def add_inventory(r: Report):
	r.h1("10. Source and extension inventory")
	r.h2("Public Python symbols by module")
	r.p(
		"The inventory below is derived with Python AST parsing. It is a navigation aid rather than a promise that every symbol is a stable external API; consumers should prefer the services explicitly described earlier."
	)
	r.table(["Source module", "Top-level public classes/functions"], python_inventory())
	r.h2("Scheduler and framework hooks")
	r.table(
		["Hook", "Configured value"],
		[
			[
				"permission_query_conditions",
				"Wildcard → access.services.enforcement.get_permission_query_conditions",
			],
			["has_permission", "Wildcard → access.services.enforcement.has_permission"],
			["scheduler_events", "Cron */5 * * * * → article.publish_scheduled_articles"],
			[
				"fixtures",
				"Identification Type, Stakeholder Entity, Stakeholder Classification filtered seed sets",
			],
		],
	)


def add_governance(r: Report):
	r.h1("11. Governance, limitations and change management")
	r.h2("Known boundaries")
	r.bullets(
		[
			"This document describes repository source, not customisations stored only in a site's database.",
			"Satellite applications can add fields, hooks, roles and behaviours; their documentation must be read alongside this reference.",
			"The OneRC Workspace currently contains a heading but no shortcuts, charts, cards or quick lists in source.",
			"No outgoing SMS provider, queueing workflow or delivery-status integration is implemented in this repository.",
			"A number of DocType test modules are generated scaffolds rather than substantive tests.",
			"API compatibility/versioning is not explicitly declared; coordinate breaking changes with all consumers.",
		]
	)
	r.h2("Recommended change process")
	r.numbered(
		[
			"Classify the change by shared-Core versus product-domain responsibility; keep product facts in satellites.",
			"Update DocType schema, server validation, client behaviour and tests together.",
			"For access changes, test query, document and explicit guard layers plus fail-closed paths.",
			"For integration hooks, preserve inert defaults, ownership boundaries and conflict detection.",
			"Add a forward migration patch when stored data must change; never rely only on a renamed field.",
			"Regenerate this document and review its data dictionary against the intended schema.",
		]
	)
	r.h2("Glossary")
	r.table(
		["Term", "Meaning"],
		[
			["Core", "This shared onerc_core Frappe application."],
			[
				"Satellite app",
				"A separately installed product/domain app integrating with Core through declared contracts.",
			],
			["DocType", "Frappe's metadata-backed document model, form and permission unit."],
			["Single", "A one-record DocType used for site-wide settings."],
			[
				"Nested set",
				"Tree storage using left/right bounds for efficient ancestor and descendant queries.",
			],
			["Geo scope", "Set of geographic nodes on which a user may exercise a particular role."],
			["Affiliation", "Derived link between a Red Profile and a domain record/type."],
			["Capability gate", "Policy check controlling whether an affiliation row may be revealed."],
			[
				"Fail closed",
				"When policy/configuration is missing or invalid, deny rather than widen access.",
			],
		],
	)


def build():
	dts = doctypes()
	r = Report()
	r.title_page()
	r.toc()
	add_overview(r, dts)
	add_architecture(r)
	add_modules(r)
	add_workflows(r)
	add_api(r)
	add_security(r)
	add_operations(r)
	add_testing(r)
	add_data_dictionary(r, dts)
	add_inventory(r)
	add_governance(r)
	r.save()
	print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
	build()
